from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
XLSX_PATH = Path(r"D:\datasets\mi-prawo-jazdy-2026\gov-source\2026-04-08\katalog_dla_kandydatow_na_kierowcow_2026.xlsx")
OUTPUT_DIR = ROOT / "output" / "analysis" / "remaining-explanations" / "pt-only-manual-review"
TARGET_IDS = {"2649", "4146", "4104", "4105"}
MANUAL_NOTES = {
    "2649": "Oficjalny XLSX wskazuje odpowiedź A = 'Gołoledź' na pytanie o czynnik zmniejszający możliwość poślizgu. To wygląda na merytoryczną sprzeczność w samym źródle i wymaga decyzji redakcyjnej albo weryfikacji z ekspertem.",
    "4104": "Oficjalny XLSX zawiera pytanie o ten sam prompt i ten sam obraz co 4105, ale z inną poprawną odpowiedzią. Konflikt jest źródłowy, nie powstał w naszym imporcie.",
    "4105": "Oficjalny XLSX zawiera pytanie o ten sam prompt i ten sam obraz co 4104, ale z inną poprawną odpowiedzią. Konflikt jest źródłowy, nie powstał w naszym imporcie.",
    "4146": "Oficjalny XLSX wskazuje odpowiedź C = 'Większe niż wartość określona dla danego typu tramwaju', co wygląda technicznie podejrzanie. To wymaga weryfikacji z dokumentacją tramwajową albo ekspertem.",
}


def load_source_rows() -> list[dict[str, str]]:
    workbook = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    sheet = workbook["katalog"]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    header_map = {str(value).strip(): index for index, value in enumerate(header)}

    rows: list[dict[str, str]] = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        external_id = row[header_map["Numer pytania"]]
        if external_id is None:
            continue

        external_id = str(external_id).strip()
        if external_id not in TARGET_IDS:
            continue

        rows.append({key: ("" if row[index] is None else str(row[index])) for key, index in header_map.items()})

    rows.sort(key=lambda item: int(item["Numer pytania"]))
    return rows


def build_report_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    report_rows = []

    for row in rows:
        external_id = row["Numer pytania"]
        report_rows.append(
            {
                "external_id": external_id,
                "prompt": row["Pytanie"],
                "correct_answer": row["Poprawna odp"],
                "option_a": row["Odpowiedź A"],
                "option_b": row["Odpowiedź B"],
                "option_c": row["Odpowiedź C"],
                "media": row["Media"],
                "scope": row["Zakres struktury"],
                "note": MANUAL_NOTES[external_id],
            }
        )

    return report_rows


def write_markdown(rows: list[dict[str, str]], output_path: Path) -> None:
    lines = ["# PT Source-Verified Anomalies", ""]

    for row in rows:
        lines.extend(
            [
                f"## {row['external_id']}",
                "",
                f"- Pytanie: {row['prompt']}",
                f"- Poprawna odpowiedź w XLSX: {row['correct_answer']}",
                f"- A: {row['option_a']}",
                f"- B: {row['option_b']}",
                f"- C: {row['option_c']}",
                f"- Media: {row['media'] or '(brak)'}",
                f"- Zakres: {row['scope']}",
                f"- Uwaga: {row['note']}",
                "",
            ]
        )

    output_path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    source_rows = load_source_rows()
    report_rows = build_report_rows(source_rows)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    json_path = OUTPUT_DIR / "source-verified-anomalies.json"
    md_path = OUTPUT_DIR / "source-verified-anomalies.md"

    json_path.write_text(json.dumps(report_rows, ensure_ascii=False, indent=2), encoding="utf-8")
    write_markdown(report_rows, md_path)

    print(
        json.dumps(
            {
                "anomaly_count": len(report_rows),
                "json_path": str(json_path),
                "markdown_path": str(md_path),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
